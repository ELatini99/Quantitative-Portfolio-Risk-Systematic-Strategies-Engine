from __future__ import annotations

from pathlib import Path
from typing import Mapping

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

NAVY = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREY = "F2F2F2"
WHITE = "FFFFFF"
GREY = "B7C9D6"


def write_table(df: pd.DataFrame, path: Path, *, index: bool = False,
                float_format: str = "%.6f") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=index, encoding="utf-8-sig", float_format=float_format)


def add_metadata(df: pd.DataFrame, *, as_of: str, description: str,
                 units: dict[str, str] | None = None) -> pd.DataFrame:
    out = df.copy()
    out.insert(0, "As Of", as_of)
    out.insert(1, "Description", description)
    if units:
        for col, unit in units.items():
            if col in out.columns:
                out[f"{col} Unit"] = unit
    return out


def _format_sheet(ws, *, freeze="A4", percent_columns=None,
                  decimal_columns=None, integer_columns=None,
                  conditional_matrix=False):
    percent_columns = percent_columns or set()
    decimal_columns = decimal_columns or set()
    integer_columns = integer_columns or set()

    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    thin = Side(style="thin", color=GREY)

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center")

    headers = {cell.value: cell.column for cell in ws[1]}
    for name, col in headers.items():
        letter = get_column_letter(col)
        for r in range(2, ws.max_row + 1):
            c = ws.cell(r, col)
            if name in percent_columns:
                c.number_format = "0.00%"
            elif name in decimal_columns:
                c.number_format = "0.0000"
            elif name in integer_columns:
                c.number_format = "#,##0"

    for col in range(1, ws.max_column + 1):
        max_len = max(len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 100) + 1))
        ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 11), 28)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 30

    if conditional_matrix and ws.max_row >= 2 and ws.max_column >= 2:
        rng = f"B2:{get_column_letter(ws.max_column)}{ws.max_row}"
        ws.conditional_formatting.add(
            rng,
            ColorScaleRule(start_type="num", start_value=-1, start_color="F4CCCC",
                           mid_type="num", mid_value=0, mid_color="FFFFFF",
                           end_type="num", end_value=1, end_color="C9DAF8"),
        )


def _add_title(ws, title: str, subtitle: str | None = None):
    ws.insert_rows(1, amount=3 if subtitle else 2)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ws.max_column, 2))
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color=NAVY)
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(ws.max_column, 2))
        ws["A2"] = subtitle
        ws["A2"].font = Font(size=10, italic=True, color="666666")


def create_excel_report(path: Path, tables: Mapping[str, pd.DataFrame],
                        titles: Mapping[str, str], subtitles: Mapping[str, str] | None = None,
                        matrix_sheets: set[str] | None = None,
                        percent_columns: Mapping[str, set[str]] | None = None,
                        decimal_columns: Mapping[str, set[str]] | None = None,
                        integer_columns: Mapping[str, set[str]] | None = None,
                        figures: Mapping[str, Path] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    subtitles = subtitles or {}
    matrix_sheets = matrix_sheets or set()
    percent_columns = percent_columns or {}
    decimal_columns = decimal_columns or {}
    integer_columns = integer_columns or {}

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet, df in tables.items():
            df.to_excel(writer, sheet_name=sheet[:31], index=False)

    wb = load_workbook(path)
    for sheet, df in tables.items():
        ws = wb[sheet[:31]]
        _add_title(ws, titles.get(sheet, sheet), subtitles.get(sheet))
        header_row = 4 if sheet in titles and subtitles.get(sheet) else 3
        _format_sheet(
            ws,
            freeze=f"A{header_row + 1}",
            percent_columns=percent_columns.get(sheet, set()),
            decimal_columns=decimal_columns.get(sheet, set()),
            integer_columns=integer_columns.get(sheet, set()),
            conditional_matrix=sheet in matrix_sheets,
        )
        if sheet in {"SUMMARY", "RISK ANALYTICS", "STRESS TEST"}:
            ws.sheet_view.showGridLines = False

    if figures:
        for sheet, img_path in figures.items():
            if sheet[:31] not in wb.sheetnames or not img_path.exists():
                continue
            ws = wb[sheet[:31]]
            img = XLImage(str(img_path))
            img.width = 720
            img.height = 400
            ws.add_image(img, f"A{ws.max_row + 3}")

    wb.save(path)
